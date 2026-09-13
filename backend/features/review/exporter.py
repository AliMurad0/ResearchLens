"""
CONCEPT: We use ReportLab for PDF generation (not fpdf2) -- ReportLab's
Paragraph objects handle text wrapping automatically, including
breaking long unbreakable strings like DOI URLs, which avoided a class
of crash the older fpdf2 version kept hitting.

References now arrive as structured dicts (title, authors, year, venue,
doi...) rather than pre-formatted strings, since the rest of the app
needs those structured fields too. _format_reference() is the one
place that turns a structured paper back into an APA-style citation
line for the document.
"""

import re
import html
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from config import settings


def _safe_filename(topic: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", topic.strip()).strip("_").lower()


def _format_reference(ref: dict) -> str:
    authors = ref.get("authors") or ["Unknown authors"]
    author_str = ", ".join(authors)
    year = ref.get("year") or "n.d."
    title = ref.get("title") or "Untitled"
    venue = ref.get("venue")
    doi = ref.get("doi")

    entry = f"{author_str} ({year}). {title}."
    if venue:
        entry += f" {venue}."
    if doi:
        entry += f" https://doi.org/{doi}" if not str(doi).startswith("http") else f" {doi}"
    return entry


def export_to_docx(topic: str, review_text: str, references: list[dict]) -> str:
    doc = Document()
    doc.add_heading(f"Literature Review: {topic}", level=1)

    for paragraph in review_text.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())

    doc.add_heading("References", level=2)
    for ref in references:
        doc.add_paragraph(_format_reference(ref), style="List Number")

    settings.outputs_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{_safe_filename(topic)}_review.docx"
    doc.save(settings.outputs_dir / filename)
    return filename


def export_to_pdf(topic: str, review_text: str, references: list[dict]) -> str:
    settings.outputs_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{_safe_filename(topic)}_review.pdf"
    filepath = settings.outputs_dir / filename

    doc = SimpleDocTemplate(
        str(filepath), pagesize=A4,
        topMargin=20 * mm, bottomMargin=20 * mm,
        leftMargin=20 * mm, rightMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    ref_style = ParagraphStyle("ref", parent=styles["BodyText"], fontSize=9, leading=12)

    story = [Paragraph(html.escape(f"Literature Review: {topic}"), styles["Heading1"]), Spacer(1, 12)]

    for paragraph in review_text.split("\n\n"):
        if paragraph.strip():
            story.append(Paragraph(html.escape(paragraph.strip()), styles["BodyText"]))
            story.append(Spacer(1, 8))

    story.append(Spacer(1, 6))
    story.append(Paragraph("References", styles["Heading2"]))
    for i, ref in enumerate(references, 1):
        story.append(Paragraph(f"{i}. {html.escape(_format_reference(ref))}", ref_style))

    doc.build(story)
    return filename
